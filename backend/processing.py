"""Document reading, cleaning and chunking."""
import re
from .config import CHUNK_SIZE, CHUNK_OVERLAP


def clean_text(t: str) -> str:
    t = t.replace("\x00", " ")
    t = re.sub(r"-\n(\w)", r"\1", t)        # join hyphenated line breaks
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def extract_pdf(path) -> list[dict]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return [{"page": i + 1, "section": None, "text": clean_text(p.extract_text() or "")}
            for i, p in enumerate(reader.pages)]


def extract_docx(path) -> list[dict]:
    """DOCX has no fixed pages, so we cite the nearest heading as the 'section'."""
    from docx import Document
    doc = Document(str(path))
    sections, current, buf = [], "Introduction", []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if (p.style.name or "").lower().startswith(("heading", "title")):
            if buf:
                sections.append({"page": None, "section": current, "text": clean_text("\n\n".join(buf))})
            current, buf = txt, []
        else:
            buf.append(txt)
    for table in doc.tables:                      # include table text too
        for row in table.rows:
            buf.append(" | ".join(c.text.strip() for c in row.cells))
    if buf:
        sections.append({"page": None, "section": current, "text": clean_text("\n\n".join(buf))})
    return sections


def extract_document(path, ext: str) -> list[dict]:
    return extract_pdf(path) if ext == "pdf" else extract_docx(path)


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = text.strip()
    if not text:
        return []
    units = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= size:
            units.append(para)
        else:
            units += [s.strip() for s in re.split(r"(?<=[.!?])\s+", para) if s.strip()]
    fixed = []
    for u in units:                                # hard-split very long sentences
        while len(u) > size:
            fixed.append(u[:size])
            u = u[size - overlap:]
        fixed.append(u)
    chunks, cur = [], ""
    for u in fixed:
        if cur and len(cur) + len(u) + 1 > size:
            chunks.append(cur)
            tail = cur[-overlap:] if overlap else ""
            tail = tail[tail.find(" ") + 1:] if " " in tail else tail
            cur = (tail + " " + u).strip()
        else:
            cur = (cur + " " + u).strip() if cur else u
    if cur:
        chunks.append(cur)
    return chunks


def build_chunks(sections: list[dict]) -> list[dict]:
    out = []
    for s in sections:
        for c in chunk_text(s["text"]):
            out.append({"text": c, "page": s["page"], "section": s["section"]})
    return out
